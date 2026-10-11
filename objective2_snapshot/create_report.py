"""Generate the extended assessment report from verified, computed Python outputs."""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
P=Path(__file__).resolve().parent; O=P/'output'; D=P/'data'
r1=json.loads((O/'model_2_1.json').read_text());r2=json.loads((O/'model_2_2.json').read_text())
doc=Document(); sec=doc.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21);sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.6);sec.left_margin=Cm(1.75);sec.right_margin=Cm(1.75)
for n in ['Normal','Body Text','Caption']:
 st=doc.styles[n];st.font.name='Calibri';st.font.size=Pt(10);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1
for n,size in [('Title',22),('Heading 1',16),('Heading 2',12)]:
 st=doc.styles[n];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(8);st.paragraph_format.space_before=Pt(8)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('HIT140 Objective 2  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
for r in foot.runs:r.font.size=Pt(8)
def p(t):return doc.add_paragraph(t)
def h(t):doc.add_heading(t,1)
def sub(t):doc.add_heading(t,2)
def page(t):
 h(t);doc.paragraphs[-1].paragraph_format.page_break_before=True
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,s in zip(t.rows[0].cells,headers):c.text=str(s)
 for values in rows:
  for c,v in zip(t.add_row().cells,values):c.text=str(v)
 pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
 pr.append(b)
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,cell in enumerate(row.cells):
   if widths:cell.width=Cm(widths[j])
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF1' if i==0 else ('F7F8F9' if i%2==0 else 'FFFFFF'));cp.append(shade)
   mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),'75');x.set(qn('w:type'),'dxa');mar.append(x)
   cp.append(mar)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2)
    for run in para.runs:run.font.size=Pt(9);run.bold=i==0
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

def pic(name,caption,width=17.3):
 q=doc.add_paragraph();q.paragraph_format.keep_with_next=True;r=q.add_run();r.add_picture(str(O/name),width=Cm(width))
 p(caption).style=doc.styles['Caption']
def eq(text):
 para=doc.add_paragraph();math=OxmlElement('m:oMath');run=OxmlElement('m:r');tx=OxmlElement('m:t');tx.text=text;run.append(tx);math.append(run);para._p.append(math)
def metric_table(label):
 df=pd.read_csv(O/f'metrics_{label}.csv');df=df[df.split=='Holdout']
 table(['Holdout method','MAE','RMSE','R²'],[[r.method,f'{r.MAE:.3f}',f'{r.RMSE:.3f}',f'{r.R2:.3f}'] for r in df.itertuples()],[9.6,2.6,2.6,2.6])
def coef_table(label):
 df=pd.read_csv(O/f'coefficients_{label}.csv')
 rows=[]
 for i,r in enumerate(df.itertuples()):
  rows.append(['Intercept' if i==0 else f'x{i}',f'{r.coefficient:.4f}',f'{r.lower95:.3f} to {r.upper95:.3f}','<0.001' if r.p_value<.001 else f'{r.p_value:.3f}'])
 table(['Term','Coefficient','95% confidence interval','p value'],rows,[2.4,3.5,8.0,3.5])
def cv_table(label):
 df=pd.read_csv(O/f'temporal_validation_{label}.csv');rows=[]
 for fold,g in df.groupby('fold'):
  a=g[g.method=='OLS eight predictors'].iloc[0];b=g[g.method=='Training mean baseline'].iloc[0]
  rows.append([str(fold),f'{a.train_matches} / {a.validation_matches}',f'{a.MAE:.3f}',f'{b.MAE:.3f}',f'{a.R2:.3f}'])
 table(['Fold','Train / validation matches','OLS MAE','Mean baseline MAE','OLS R²'],rows,[1.2,5.0,3.4,4.4,3.4])

# Page 1
h0=doc.add_paragraph('Predicting FIFA World Cup 2026 outcomes with linear regression','Title')
p('HIT140  |  Objective 2  |  Extended analytical report  |  5 October 2026')
sub('Purpose and main findings')
p('This study evaluates two pre-match forecasting questions. Regression 2.1 predicts the goal difference between two teams from 104 match observations. Regression 2.2 predicts a team’s goals from 208 team–match observations. Each multiple linear regression uses exactly eight explanatory variables. Four predictor concepts are shared; four are specific to each task. All wrangling, statistical analysis and figures were produced in Python.')
p('The goal-difference model achieved a held-out mean absolute error of 1.146 goals and R² of 0.128. The team-goals model achieved a mean absolute error of 0.971 goals and R² of 0.024. Both improved on predictions based on their training-set means, but the team-goals model did not outperform an Elo-only regression. These results support modest forecasting value, with stronger evidence for predicting differences than precise team scoring totals. They do not support claims of highly accurate score prediction.')
table(['Design element','Regression 2.1','Regression 2.2'],[
 ['Unit of observation','One completed match','One team in one completed match'],['Dataset rows','104','208'],['Explanatory variables','8 numeric variables','8 numeric variables'],['Outcome','Team A goals minus team B goals','Goals scored by the team'],['Training observations','85 matches','170 team rows from 85 matches'],['Held-out observations','19 matches','38 team rows from the same 19 matches']],[5.2,6.1,6.1])
sub('Outcome definitions')
p('Team A is the first-listed team in the FIFA match report; team B is the second-listed team. This ordering is fixed before inspecting the result. A positive goal difference means A outscored B, a negative difference means B outscored A, and zero means the score was level. The label A does not imply that the match was played at A’s home venue.')
p('Goals include extra time when it was played, but exclude penalty-shootout kicks. This follows the full-match score in the FIFA reports and the historical dataset definition [1, 2]. A match decided on penalties therefore has zero goal difference. The target is the eventual match score, not the 90-minute score. Actual extra-time duration is never used as a predictor because it is unknown before kickoff.')
sub('Research questions')
p('Regression 2.1 asks: how accurately can prior relative strength, results, recovery, venue advantage and historical experience predict the final goal difference? Regression 2.2 asks: how accurately can prior attacking output, the opponent’s defensive record and pre-match context predict the number of goals scored by one team?')

# Page 2
page('Data acquisition and preparation')
p('The outcome dataset is a census of all 104 matches, covering 48 teams, rather than a random sample of the tournament. FIFA’s 104 full-time match reports were downloaded and archived. Their match IDs, local match dates, teams, venue countries, final scores, extra-time flags and source URLs were extracted with Python. The reports show 308 goals, nine matches with extra time and four shootouts [1]. Historical features were constructed from Mart Jürisoo’s international-results collection [2].')
p('After harmonising team names and allowing for reversed team order, every FIFA outcome and match date agreed with the corresponding historical record. A separate fixtures webpage was considered during acquisition, but its date and match-number conventions differed; it was not used to determine the final model data. FIFA reports remain the authoritative source for tournament outcomes and dates.')
sub('Cleaning decisions')
p('Aliases were mapped explicitly, including USA to United States, Korea Republic to South Korea, Czechia to Czech Republic, IR Iran to Iran, Cabo Verde to Cape Verde, and Côte d’Ivoire to Ivory Coast. Joining used both team identities and verified match context rather than assuming resource numbers equal match numbers. The mapping is preserved in build_data.py.')
p('Only historical records dated before 11 June 2026 were retained as the initial history. Two records for Tahiti versus New Caledonia on 17 February 1974 had the same identifying fields but different scores. Both were excluded as unresolved duplicates and retained in an exclusions file. The resulting initial history contains 49,414 matches. Current-tournament history was then appended from FIFA reports, preventing duplicate inclusion of 2026 World Cup matches. No tournament row was removed or imputed.')
sub('Preventing information leakage')
p('For each match date, features were calculated before adding any result from that date to the histories. Recent form uses the latest ten completed internationals strictly before the match’s date. Every participating team had ten available prior matches, so no cold-start imputation was required. Same-day results are excluded conservatively because the history records dates rather than reliable kickoff and completion timestamps.')
p('The audit file stores each side’s ten source matches, their dates, scores and competition indicators. This makes it possible to check that neither the current outcome nor a future match entered an explanatory variable. All later tournament totals, current-match shots, possession, cards, scorers and actual extra-time duration are excluded from the design matrices. Data were retrieved after the event, so this is a retrospective backtest using historically available events, not evidence of forecasts recorded live before kickoff.')
sub('Files and integrity')
p('regression_2_1.csv contains one ID, eight explanatory columns and one outcome. regression_2_2.csv has the same column structure at team–match level. Dates, team names, source URLs and split labels are stored separately as metadata; they are not additional predictors. Source files have SHA-256 hashes, and validation checks enforce row counts, unique IDs, eight-column feature lists, finite values, correct targets and strict chronological history.')

# Page 3
page('Regression 2 1 explanatory variables')
p('The model represents relative advantage by subtracting team B’s pre-match value from team A’s value. These eight predictors were fixed on football grounds before examining held-out model performance. They cover strength, recent outcomes, rest, location and experience. No transformations were selected by searching for a favourable test score.')
table(['Term','Predictor and exact construction','Reason for inclusion'],[
 ['x1','Elo difference: (rating A − rating B) / 100.','Summarises past results while accounting for opponent strength. Expected positive association.'],
 ['x2','Win-rate difference over each side’s last 10 prior internationals.','Captures recent success that may differ from long-run strength. Expected positive association.'],
 ['x3','Draw-rate difference over the last 10 prior internationals.','Measures recent tendency toward level scores. Direction is not predetermined.'],
 ['x4','Rest difference: min(days since A last played, 14) − min(days since B last played, 14).','Represents relative recovery. Capping prevents long breaks dominating the scale.'],
 ['x5','Home advantage: +1 when playing in A’s country; −1 in B’s country; 0 otherwise.','Represents venue familiarity and local support. Expected positive association.'],
 ['x6','World Cup experience difference: log(1 + A’s prior finals matches) − log(1 + B’s).','Measures historical tournament exposure with diminishing increments. Expected positive association.'],
 ['x7','Head-to-head balance: sum of A’s win/draw/loss signs (+1/0/−1) in up to 5 prior meetings, divided by meetings + 2.','Captures matchup history. Shrinkage toward zero limits tiny-history effects; no meetings gives zero.'],
 ['x8','Competitive-share difference: A’s fraction of non-Friendly games in its last 10 minus B’s fraction.','Distinguishes preparation in competitive settings from friendly-heavy schedules. Direction is uncertain.']],[1.1,8.5,7.8])
sub('Elo construction')
p('Elo is calculated by this analysis, not presented as an official FIFA ranking. Each team starts at 1,500 before its first archived match. Expected result = 1 / (1 + 10 raised to [(opponent rating − team rating − home adjustment) / 400]). The update is 20 × (actual result − expected result), with actual result 1, 0.5 or 0. The opponent receives the opposite adjustment. A home side receives a 100-point expectation adjustment; neutral matches receive zero. Same-day adjustments are applied after features are recorded. Goal margins and shootout winners do not affect these updates.')
p('The full historical sequence begins in 1872. Historical successor-team naming follows the source. World Cup experience counts only finals matches, including completed earlier matches in 2026; qualifiers do not count. Competitive share follows the source’s tournament label and does not measure the difficulty of individual opponents.')

# Page 4
page('Regression 2 1 estimation and interpretation')
p('Ordinary least squares estimates an intercept and eight slopes by minimising the sum of squared residuals. The development model uses the 85 matches played through 2 July 2026. Its design matrix has full rank, so all nine parameters are identifiable. The fitted equation predicts D, the goal difference, using x1 to x8 in the order defined on the preceding page [3].')
eq('D = 0.3040 + 0.0228x1 + 2.5443x2 + 2.4287x3 + 0.0832x4')
eq('                 + 0.5340x5 + 0.4463x6 + 1.4453x7 − 0.6391x8')
p('Training R² is 0.506 and adjusted R² is 0.454. The reduction after adjustment reflects the cost of fitting eight predictors to a modest sample. These are in-sample measures and are not estimates of future accuracy.')
coef_table('2_1')
p('The table reports HC3 heteroscedasticity-robust standard errors and 95% confidence intervals. For each slope, the exploratory null hypothesis is that the coefficient is zero after accounting for the other seven predictors. The robust joint F statistic is 7.689, with p < 0.001. This supports some joint association in the training data, not a claim of strong out-of-sample forecasting.')
sub('Practical meaning of the estimates')
p('A 0.10 increase in A’s win-rate advantage corresponds to 0.254 additional predicted goals of difference, holding other inputs fixed. Its interval includes zero, so the data do not establish a precise independent win-rate effect. An additional capped rest day corresponds to 0.083 goals of difference, also with substantial uncertainty.')
p('The World Cup experience coefficient is 0.446 per unit of log-count difference (95% CI 0.083 to 0.810). Multiplying the ratio (1 + A’s experience)/(1 + B’s experience) by two corresponds to approximately 0.309 additional predicted goals, since 0.446 × log(2) = 0.309. This is an association that may reflect unmeasured team quality, not evidence that historical appearances cause present goals.')
p('Elo’s conditional slope is small and imprecise despite the useful Elo-only baseline. Elo and other strength indicators overlap, increasing coefficient uncertainty. The positive draw-rate slope should not be interpreted as a strategy to seek draws; it is a conditional pattern in a small observational dataset. Individual p values are exploratory and unadjusted for multiple comparisons.')

# Page 5
page('Regression 2 1 predictive evaluation')
p('The final 19 matches, played from 3 to 19 July, form the chronological holdout. No match from those dates was used to estimate the reported regression coefficients. Earlier held-out results are allowed to update subsequent pre-match features once they have occurred; coefficients remain frozen. This evaluates sequential pre-match forecasting, not a single forecast of the whole bracket on 2 July.')
metric_table('2_1')
p('MAE is the average absolute prediction error in goals. RMSE gives greater weight to large errors. R² compares squared error with variation around the held-out outcome mean; a negative value means performance worse than that test-mean reference. The fitted model’s MAE is 1.146 goals, a 23.2% reduction from the training-mean baseline. It also outperforms the Elo-only and zero-difference baselines in this holdout [4].')
p('Held-out R² is only 0.128, well below the training value of 0.506. The prediction task therefore remains difficult even though errors improve over the specified baselines. The held-out outcomes range from −3 to +2 goals. They are all knockout matches, while the training set includes the full group stage and 13 round-of-32 matches. This change in match composition limits generalisation to an arbitrary future fixture.')
sub('Development validation')
p('Three expanding-window validation folds were constructed within the development period. Whole calendar dates are kept together. Folds validate on 17–22 June, 23–27 June and 28 June–2 July respectively. The first fold has only 20 training matches for nine fitted parameters, so its estimates are particularly unstable. These folds provide a stability check rather than a search for a favourable specification [5].')
cv_table('2_1')
p('The eight-variable model performs worse than the mean baseline in the earliest fold but improves in the later folds. This inconsistency is evidence against treating one good test score as a reliable universal accuracy estimate.')
sub('Uncertainty in the apparent improvement')
p('A 3,000-resample bootstrap of whole held-out matches, with the trained coefficients fixed, gives an approximate 95% interval of 0.793 to 1.521 goals for MAE. The interval for MAE improvement over the mean baseline is −0.071 to 0.748 goals. Because it includes zero, the observed improvement is not established precisely by these 19 matches. These are conditional performance intervals, not individual-match prediction intervals; recurring teams and temporal dependence can make them overoptimistic.')

# Page 6
page('Regression 2 1 diagnostics and limitations')
p('Diagnostics were computed using training residuals, avoiding specification changes driven by held-out errors. Residual plots check patterns left unexplained by the linear model; Q-Q plots compare residual tails with a normal reference. Cook’s distance identifies observations whose removal may materially change fitted coefficients.')
pic('diagnostics_2_1.png','Figure 1. Python-generated training diagnostics and held-out predictions for Regression 2.1. The dashed Cook’s-distance line is the screening value 4/n; it is not a deletion rule.')
p('The residual-versus-fitted plot has visible bands because observed goal differences are integers. It does not establish perfect linearity. The Q-Q plot is reasonably aligned centrally with some tail departures. Shapiro–Wilk p = 0.225 and Breusch–Pagan p = 0.578 do not provide strong evidence against residual normality or constant variance, respectively; failure to reject does not prove either assumption.')
p('Maximum variance inflation factor is 8.006 for Elo difference, followed by 4.632 for win-rate difference and 3.633 for World Cup experience. This explains why a useful strength signal may have an uncertain conditional slope. The eight specified predictors are retained to preserve the task design; any future reduced or regularised model would require separate validation.')
p('Six training observations exceed Cook’s distance 4/85. The most influential is Ecuador versus Germany, match M056. Omitting this training match in a labelled sensitivity analysis changes held-out MAE from 1.146 to 1.143 and RMSE from 1.395 to 1.403. The largest individual prediction change is 0.420 goals. The main model retains this genuine match: unusual results are part of the forecasting task, not data errors.')
p('HC3 addresses unequal error variance but not dependence created by teams appearing repeatedly. Team A ordering also matters because the fitted intercept is nonzero: exchanging the labels is not guaranteed to negate the prediction exactly. All evaluation therefore uses the fixed FIFA ordering. A future model could impose antisymmetry, but that is not the fitted specification assessed here.')

# Page 7
page('Regression 2 2 explanatory variables')
p('Each match is expanded into two rows, one for each team. The outcome is that team’s goals. Opponent information is joined from the opposing team’s pre-match history, never from the target match itself. The two rows are dependent observations and remain together throughout evaluation.')
table(['Term','Predictor and construction','Expected role'],[
 ['x1','Elo difference: (team rating − opponent rating) / 100.','Relative quality; expected positive slope.'],
 ['x2','Rest difference: capped team rest days − capped opponent rest days, with the same 14-day cap.','Relative recovery; an uncertain but potentially positive effect.'],
 ['x3','Home advantage: +1 in the team’s country, −1 in the opponent’s country, otherwise 0.','Venue context and support; expected positive slope.'],
 ['x4','Competitive share: fraction of the team’s last 10 prior games labelled other than Friendly.','Context of preparation; expected sign is uncertain.'],
 ['x5','Goals-for mean: team’s mean goals scored across its last 10 prior internationals.','Recent attacking output; expected positive slope.'],
 ['x6','Opponent goals-against mean: opponent’s mean goals conceded across its last 10 prior internationals.','Opposition defensive vulnerability; expected positive slope.'],
 ['x7','Failed-to-score rate: fraction of team’s last 10 prior internationals with zero goals scored.','Frequency of attacking failure; expected negative slope.'],
 ['x8','Opponent clean-sheet rate: fraction of opponent’s last 10 prior internationals with zero goals conceded.','Frequency of defensive shutouts; expected negative slope.']],[1.1,10.0,6.3])
sub('Checking the overlap rule')
p('The models share exactly four underlying predictor concepts: Elo difference, rest difference, home advantage and competitive share. Competitive share is conservatively counted as shared even though Regression 2.1 uses a difference and Regression 2.2 uses the team’s level. Renaming or re-expressing a variable is not treated as making it new.')
p('Regression 2.1 uniquely uses win rate, draw rate, World Cup experience and head-to-head balance. Regression 2.2 uniquely uses mean goals scored, opponent mean goals conceded, failure-to-score rate and opponent clean-sheet rate. These are different summaries and forecasting mechanisms, although some naturally use the same historical score records. Goals means describe magnitude; zero-goal frequencies describe how often attacks fail or defences shut out opponents.')
sub('Descriptive context')
p('Across all 208 rows, team goals average 1.481 with sample SD 1.393, median 1 and range 0–7. The team’s prior scoring mean averages 1.877, while the opponent’s prior concession mean averages 0.971. These differences reflect different populations of opponents and fixtures; they are not quantities expected to equal tournament scoring. Elo difference, rest difference and home advantage each average zero across the paired dataset by construction.')

# Page 8
page('Regression 2 2 estimation and interpretation')
p('The team-goals model uses the same 85 development matches, yielding 170 rows. Ordinary least squares again estimates eight slopes and an intercept. Goals are discrete and nonnegative, but linear regression is fitted as required by the task. Its prediction G represents an estimated conditional mean; it need not be an integer.')
eq('G = 0.3615 + 0.3597x1 + 0.0480x2 + 0.7755x3 + 1.0288x4')
eq('                      + 0.1915x5 + 0.2211x6 + 0.2790x7 − 0.1589x8')
p('Training R² is 0.314 and adjusted R² is 0.280. The confidence intervals below use standard errors clustered by match, allowing the two team residuals within the same fixture to be correlated. This changes the uncertainty calculation, not the OLS coefficient estimates. There are 85 training match clusters.')
coef_table('2_2')
p('The robust joint F statistic is 9.095 with p < 0.001. A 100-point Elo advantage corresponds to 0.360 additional predicted goals, holding the other predictors constant (95% CI 0.193 to 0.526). This is the clearest estimated independent association in the model.')
p('Playing in the team’s own country rather than at a neutral venue corresponds to 0.775 more predicted goals (95% CI 0.263 to 1.288). Because the coding is −1, 0, +1, moving from the opponent’s home country to the team’s own country spans two units. This coefficient is supported by few host-team fixtures and combines location with omitted contextual factors; it should not be interpreted as a universal causal home effect.')
p('An increase of one in the team’s prior average goals corresponds to 0.192 more predicted goals, but its confidence interval spans −0.310 to 0.693. The prior zero-scoring rate unexpectedly has a positive fitted slope. Its wide interval includes both signs, so there is no evidence that failing to score helps future attacks. Correlated attack summaries and limited data can reverse conditional coefficients.')
p('Proportion predictors must be interpreted on their actual scale: a ten-percentage-point change equals 0.10 predictor units. P values are exploratory and unadjusted. Match clustering does not account fully for dependence between separate matches involving the same teams.')

# Page 9
page('Regression 2 2 predictive evaluation')
p('Evaluation uses 38 team rows from the same 19 held-out matches as Regression 2.1. Splitting individual team rows randomly would allow one side of a match into training and its opponent into testing; this would misrepresent performance on unseen matches. The chronological match-level split prevents that problem.')
metric_table('2_2')
p('The eight-predictor model’s MAE is 0.971 goals, compared with 1.075 for the training-mean baseline: an improvement of 9.7%. Its RMSE improves by only 1.4%, from 1.354 to 1.335. Held-out R² is 0.024, meaning little of the held-out variation is captured. The Elo-only regression performs better, with MAE 0.946 and RMSE 1.316. Additional explanatory variables therefore do not demonstrate added predictive value in this holdout.')
sub('Development validation')
cv_table('2_2')
p('The full model beats the mean baseline in all three development folds, but those estimates do not guarantee later performance. The final holdout includes high-scoring observations and a different mix of strong teams. The 19-match holdout is small; its 38 rows do not constitute 38 independent fixtures.')
sub('Uncertainty and practical errors')
p('Resampling entire matches preserves the pairing of team rows. Across 3,000 bootstrap replicates, the approximate conditional 95% MAE interval is 0.692–1.314 goals. The interval for improvement over the training-mean baseline is −0.019 to 0.224 goals, which includes zero. Repeated teams across fixtures and the fixed-coefficient bootstrap limit the scope of this interval.')
p('In the third-place match, England scored six goals but the model predicted 1.354. In the final, it predicted 1.488 for Spain and 1.643 for Argentina, while the observed goals were one and zero. These examples show regression toward typical scoring levels and difficulty capturing extreme outcomes or individual-match events. They are taken directly from the saved holdout predictions, not chosen to demonstrate only successes.')
p('All 38 held-out predictions are nonnegative. A separately labelled sensitivity check clipping predictions at zero therefore leaves the holdout metrics unchanged. Scores are not rounded before evaluation: rounding would discard information and change the quantity being assessed. Linear regression still permits negative predictions elsewhere and supplies no natural count distribution.')

# Page 10
page('Regression 2 2 diagnostics and limitations')
pic('diagnostics_2_2.png','Figure 2. Python-generated training diagnostics and held-out predictions for Regression 2.2. Diagonal residual bands arise because goals are integer counts.')
p('The residual plot and Q-Q plot show asymmetry and departures in the tails. Breusch–Pagan p < 0.001 indicates evidence of unequal residual variance, and Shapiro–Wilk p = 0.008 indicates departures from normality. These tests are descriptive diagnostics because within-match and repeated-team dependence complicate their nominal reference distributions. The findings are consistent with the limitations of a constant-variance Gaussian error model for goal counts.')
p('Clustered standard errors improve coefficient uncertainty reporting for the paired rows but do not make residuals normal, prevent negative predictions, or improve predictive accuracy. Classical normal-theory individual prediction intervals are therefore not reported as calibrated intervals. A Poisson or negative-binomial model would be a useful separately assessed extension, but neither replaces the linear regression required here.')
p('Maximum VIF is 3.307, so multicollinearity is less severe than in Regression 2.1, although scoring means and zero-goal frequencies share information. Eight training rows exceed Cook’s-distance screening threshold 4/170. The largest is Canada’s six-goal performance against Qatar in match M027. This is a verified outcome and remains in the primary analysis.')
p('Removing the whole influential training match, including both team rows, changes held-out MAE from 0.971 to 0.960 and RMSE from 1.335 to 1.319. The maximum prediction change is 0.176 goals. This is a sensitivity analysis, not a replacement selected for better test performance. Its limited impact does not resolve the main issue: even the primary model provides little improvement over simple forecasts.')
p('Unmodelled red cards, finishing randomness, injuries and tactical decisions help explain uncertainty. Current-match events cannot be added as predictors without violating the pre-match requirement. Pre-match availability, lineups and other archived context could be considered in a later study if they can be collected consistently for every match.')

# Page 11
page('Comparison and conclusions')
table(['Assessment dimension','Regression 2.1','Regression 2.2'],[
 ['Held-out MAE','1.146 goal difference','0.971 team goals'],['Held-out RMSE','1.395','1.335'],['Held-out R²','0.128','0.024'],['Training adjusted R²','0.454','0.280'],['Improvement over mean MAE','23.2%','9.7%'],['Compared with Elo only','Lower MAE and RMSE','Higher MAE and RMSE'],['Main diagnostic concern','Overlapping strength indicators; small sample','Count outcome, unequal variance and extreme scores'],['Strongest defensible conclusion','Useful relative to tested baselines, with wide uncertainty','Limited additional value beyond a simpler strength model']],[6.0,5.7,5.7])
p('The numerical MAEs should not be compared as if the models predict the same target. A goal difference and one team’s goals have different distributions. The better comparison is each model against baselines for its own target, supported by held-out R² and uncertainty. Regression 2.1 provides the stronger evidence of improvement under this evaluation design.')
sub('Limits on generalisation')
p('The dataset covers the entire tournament, so its descriptive statistics have no sampling error for that particular finite population. Forecasting uncertainty remains because the fitted models aim to represent a broader process of football outcomes. Training, validation and testing are chronological partitions for prediction, not random samples for estimating the tournament mean. Confidence intervals rely on modelling or resampling assumptions and cannot be justified merely by having all 104 matches.')
p('The final holdout consists entirely of knockout fixtures and contains only 19 matches. Teams recur, opponents are selected by earlier success, and the variance of match outcomes can differ between stages. Historical friendlies and competitive games also differ in opposition strength and motivation. Competitive share only partially captures this. Date-level rest ignores kickoff times, travel and actual minutes played; the 14-day cap is a design choice, not an estimated biological threshold.')
p('Extra-time inclusion matches the stated full-match target but gives some knockout matches longer exposure in which to score. Actual extra time cannot be known before kickoff. The model therefore predicts a mixture of 90- and 120-minute scoring opportunities. Archived event data may also contain later corrections; no complete archive of what each website displayed at each historical kickoff was available.')
p('The two regressions are fitted separately and need not produce coherent joint score forecasts. For the final, Regression 2.1 predicts Spain’s goal difference as +0.443, while the difference between the two team-goals predictions is −0.155. Neither is overwritten to force agreement. They answer related but distinct assessment questions with different predictors and loss structures.')
sub('Final conclusion')
p('Both tasks are completed with verified FIFA outcomes, the required observation counts, exactly eight predictors each, four shared concepts, auditable pre-match histories and Python evaluation. The results support cautious use of the goal-difference model relative to the tested baselines. The team-goals model illustrates that plausible extra variables do not necessarily improve prediction. A defensible assessment reports this result and its uncertainty rather than treating a statistically significant training fit as proof of accurate forecasting.')

# Page 12
page('Reproducibility and supporting evidence')
p('The accompanying project includes raw source files, the two final modelling datasets, metadata, a data dictionary, Python scripts, all computed tables and figures, and the model predictions. No external account, API key or spreadsheet calculation is needed to reproduce the archived analysis.')
table(['File or folder','Purpose'],[
 ['raw/fifa/','104 primary-source PDF match reports and extracted text; source resource numbers are not match IDs.'],['raw/history.csv','Archived international results used to construct initial histories. Only pre-tournament rows enter features.'],['build_data.py','Extracts FIFA outcomes, standardises names, audits scores, constructs strictly prior-date features and exports datasets.'],['analyse.py','Fits both OLS models and baselines; generates temporal validation, uncertainty, diagnostics and plots.'],['validate.py','Checks row counts, feature counts, targets, pairing, prior-date inputs, split integrity and recalculated errors.'],['create_report.py','Produces this report from the recorded results and figures.'],['data/','Two model-ready CSVs plus metadata, complete history audit, source comparison, exclusions and source hashes.'],['output/','Coefficients, VIFs, descriptions, prediction errors, sensitivity results, validation tables, figures and this report.']],[5.0,12.4])
sub('Reproduction sequence')
p('Install the dependencies in requirements.txt into a Python environment, then run build_data.py, analyse.py and validate.py in that order. Run create_report.py to rebuild the Word document. collect_fifa.py can refresh the public match reports; archived snapshots should remain the reference for reproducing these exact results. The deterministic bootstrap seed is 1402026. All relative file paths are resolved from the script directory.')
p('The saved coefficients retain the evaluated training fit. All 104 and 208 observations remain in the datasets. A deployment fit to all rows would need to be labelled separately.')
sub('Sources and methods documentation')
p('[1] FIFA. FIFA World Cup 2026 full-time match reports, 104 archived PDFs, accessed 5 October 2026. The exact URL for each match is recorded in data/matches_fifa.csv. Example opening-match report: https://fdp.fifa.org/assetspublic/ce281/r12452/pdf/FullTimeMatchReport-English.pdf. Final: https://fdp.fifa.org/assetspublic/ce281/r12552/pdf/FullTimeMatchReport-English.pdf.')
p('[2] Jürisoo, M. International football results collection and data definitions. https://github.com/martj42/international_results. Raw snapshot: https://raw.githubusercontent.com/martj42/international_results/master/results.csv. Accessed 5 October 2026. Source README documents scores including extra time and excluding shootouts.')
p('[3] statsmodels. Ordinary Least Squares documentation. https://www.statsmodels.org/stable/generated/statsmodels.regression.linear_model.OLS.html. Accessed 5 October 2026. Python implementation explicitly includes an intercept.')
p('[4] scikit-learn. Regression metrics documentation. https://scikit-learn.org/stable/modules/model_evaluation.html. Accessed 5 October 2026. Metrics here are computed directly from saved predictions in Python.')
p('[5] scikit-learn. TimeSeriesSplit documentation. https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.TimeSeriesSplit.html. Accessed 5 October 2026. This project implements expanding windows on whole calendar dates and matches; it does not apply the row-based splitter directly to paired team records.')

# Page 13
page('Descriptive statistics for the final datasets')
p('These summaries describe the complete tournament datasets. They were not used to estimate preprocessing parameters or choose predictors from held-out outcomes. SD is the sample standard deviation; all proportions range from zero to one before taking differences. Definitions and scales correspond to the variable tables above.')
for label,title in [('2_1','Regression 2 1 summary'),('2_2','Regression 2 2 summary')]:
 sub(title);df=pd.read_csv(O/f'descriptive_{label}.csv');rows=[]
 for i,r in enumerate(df.itertuples(index=False)):
  rows.append([f'x{i+1}' if i<8 else 'Outcome',f'{r.mean:.3f}',f'{r.std:.3f}',f'{r.min:.3f}',f'{r[6]:.3f}',f'{r.max:.3f}'])
 table(['Variable','Mean','SD','Minimum','Median','Maximum'],rows,[2.4,3,3,3,3,3])
p('The model 2.1 outcome mean is positive because FIFA’s first-listed teams outscored their opponents more often in aggregate; it is not evidence that being printed first causes a scoring advantage. Paired predictors in model 2.2 have zero mean by algebraic construction. Full correlation matrices, individual-match prediction tables, confidence intervals, all validation-fold metrics and the VIF tables are supplied as machine-readable Python outputs.')

# Remove inherited title borders from all paragraph styles and document paragraphs.
for root in [doc.styles.element,doc.element]:
 for node in list(root.xpath('//w:pBdr')):node.getparent().remove(node)
# Add figure accessibility text.
for i,node in enumerate(doc.element.xpath('//wp:docPr')):node.set('descr',f'Regression diagnostics figure {i+1}: residual versus fitted, normal Q-Q, held-out observed versus predicted, and Cook distance panels.')
doc.core_properties.title='Predicting FIFA World Cup 2026 outcomes with linear regression'
doc.core_properties.subject='HIT140 Objective 2 extended report'
doc.core_properties.author=''
doc.save(O/'Objective_2_Extended_Report.docx')
print(O/'Objective_2_Extended_Report.docx')
